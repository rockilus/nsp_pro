"""Parse an import Excel workbook into raw row dicts.

The expected workbook has three sheets:
- ``members``  — workers with their attributes
- ``shifts``   — shift definitions
- ``schedule`` — grid: worker names in column A, dates in row 1 headers,
   shift codes in cells

Returns raw ``(members, shifts, schedule, errors)`` that the
``ImportService`` then transforms into preview DTOs.
"""

import logging
from datetime import date, datetime, time
from io import BytesIO
from typing import Any, Dict, List, Tuple

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

logger = logging.getLogger(__name__)

# ── Public API ────────────────────────────────────────────────────────────────


def parse_import_workbook(
    file_contents: bytes,
) -> Tuple[List[Dict[str, Any]], List[Dict[str, Any]], List[Dict[str, Any]], List[str]]:
    """Parse an Excel workbook and return raw dicts for each sheet.

    Returns:
        (members, shifts, schedule, errors) — each element is a list of
        row dicts; ``errors`` contains fatal problems (missing sheets, etc.).
    """
    errors: List[str] = []

    try:
        wb = load_workbook(BytesIO(file_contents), data_only=True)
    except Exception as exc:
        return [], [], [], [f"Cannot open file: {exc}"]

    # ── Validate sheets ───────────────────────────────────────────────────
    required = {"members", "shifts", "schedule"}
    available = set(wb.sheetnames)
    missing = required - available
    if missing:
        errors.append(
            f"Missing required sheet(s): {', '.join(sorted(missing))}. "
            f"Found: {', '.join(sorted(available))}"
        )
    if errors:
        return [], [], [], errors

    # ── Parse each sheet ───────────────────────────────────────────────────
    members = _parse_members_sheet(wb["members"], errors)
    shifts = _parse_shifts_sheet(wb["shifts"], errors)
    schedule = _parse_schedule_sheet(wb["schedule"], errors)

    return members, shifts, schedule, errors


# ── Sheet parsers ─────────────────────────────────────────────────────────────


def _parse_members_sheet(ws: Worksheet, errors: List[str]) -> List[Dict[str, Any]]:
    """Parse the 'members' sheet.

    Expected columns (row 1 = header):
      A: name          (required)
      B: code          (acronym — optional)
      C: start         (employment start date — required)
      D: end           (employment end date — optional)
      E: skills        (semicolon-separated — optional)
      F: contract      (weekly contract hours — optional, default 40)
      G: desired       (weekly desired hours — optional, default = contract)
      H: duty/month    (duties per month — optional, default 4)
      I: leave         (annual leave days — optional, default 20)
    """
    rows: List[Dict[str, Any]] = []
    if ws.max_row < 2:
        errors.append(
            "'members' sheet has no data rows (need at least a header + 1 row)"
        )
        return rows

    for row_idx in range(2, ws.max_row + 1):
        name = _cell_str(ws, row_idx, 1)
        if not name:
            continue  # skip empty rows

        code = _cell_str(ws, row_idx, 2)
        start = _cell_date(ws, row_idx, 3)
        end = _cell_date(ws, row_idx, 4)
        skills = _cell_str(ws, row_idx, 5)
        contract = _cell_int(ws, row_idx, 6)
        desired = _cell_int(ws, row_idx, 7)
        duty_per_month = _cell_int(ws, row_idx, 8)
        annual_leave = _cell_int(ws, row_idx, 9)

        rows.append(
            {
                "name": name.strip(),
                "code": (code or "").strip(),
                "start": start,
                "end": end,
                "skills": [s.strip() for s in (skills or "").split(";") if s.strip()],
                "contract": contract or 40,
                "desired": desired or contract or 40,
                "duty_per_month": duty_per_month or 4,
                "annual_leave": annual_leave or 20,
            }
        )

    return rows


def _parse_shifts_sheet(ws: Worksheet, errors: List[str]) -> List[Dict[str, Any]]:
    """Parse the 'shifts' sheet.

    Expected columns (row 1 = header):
      A: name             (required)
      B: code             (acronym — required)
      C: duty             (truthy/falsy — optional, default False)
      D: mandatory_rest   (truthy/falsy — optional, default False)
      E: start            (start time — required, HH:MM or Excel serial)
      F: end              (end time — required)
      G: staffing         (comma-separated — optional)
    """
    rows: List[Dict[str, Any]] = []
    if ws.max_row < 2:
        errors.append(
            "'shifts' sheet has no data rows (need at least a header + 1 row)"
        )
        return rows

    for row_idx in range(2, ws.max_row + 1):
        name = _cell_str(ws, row_idx, 1)
        if not name:
            continue

        code = _cell_str(ws, row_idx, 2)
        duty = _cell_bool(ws, row_idx, 3)
        mandatory_rest = _cell_bool(ws, row_idx, 4)
        start_time = _cell_time_minutes(ws, row_idx, 5)
        end_time = _cell_time_minutes(ws, row_idx, 6)
        staffing_raw = _cell_str(ws, row_idx, 7)

        rows.append(
            {
                "name": name.strip(),
                "code": (code or "").strip(),
                "duty": duty,
                "mandatory_rest": mandatory_rest,
                "start_time": start_time,
                "end_time": end_time,
                "staffing": [
                    s.strip() for s in (staffing_raw or "").split(",") if s.strip()
                ],
            }
        )

    return rows


def _parse_schedule_sheet(ws: Worksheet, errors: List[str]) -> List[Dict[str, Any]]:
    """Parse the 'schedule' sheet.

    Row 1: date headers (columns B onward)
    Column A: worker full names (rows 2 onward)
    Cells: shift codes, possibly multiple separated by ``" / "``,
           or the literal string ``"leave"`` (case-insensitive).
    """
    rows: List[Dict[str, Any]] = []
    if ws.max_row < 2 or ws.max_column < 2:
        errors.append("'schedule' sheet is empty or has no date columns")
        return rows

    # ── Parse date headers ────────────────────────────────────────────────
    dates: List[date] = []
    for col_idx in range(2, ws.max_column + 1):
        d = _cell_date(ws, 1, col_idx)
        if d is None:
            break
        dates.append(d)

    if not dates:
        errors.append("'schedule' sheet has no valid date headers in row 1")
        return rows

    # ── Parse worker rows ──────────────────────────────────────────────────
    for row_idx in range(2, ws.max_row + 1):
        worker_name = _cell_str(ws, row_idx, 1)
        if not worker_name:
            continue

        cells: Dict[str, List[str]] = {}
        for col_idx, d in enumerate(dates, start=2):
            raw = _cell_str(ws, row_idx, col_idx)
            if raw:
                # Split by " / " for multi-shift days
                codes = [c.strip() for c in raw.split(" / ") if c.strip()]
                if codes:
                    cells[d.isoformat()] = codes

        rows.append(
            {
                "worker_name": worker_name.strip(),
                "cells": cells,
            }
        )

    return rows


# ── Cell helpers ──────────────────────────────────────────────────────────────


def _cell_str(ws: Worksheet, row: int, col: int) -> str | None:
    """Read a cell value as a trimmed string, or None."""
    val = ws.cell(row=row, column=col).value
    if val is None:
        return None
    s = str(val).strip()
    return s if s else None


def _cell_int(ws: Worksheet, row: int, col: int) -> int | None:
    """Read a cell value as int, or None."""
    val = ws.cell(row=row, column=col).value
    if val is None:
        return None
    try:
        return int(float(str(val)))
    except ValueError, TypeError:
        logger.warning(
            "Cell %s%d is not an integer: %r", get_column_letter(col), row, val
        )
        return None


def _cell_date(ws: Worksheet, row: int, col: int) -> date | None:
    """Read a cell value as a ``date``, or None.

    Handles:
    - Python ``date`` / ``datetime`` objects (from openpyxl)
    - Excel serial date numbers (float/int)
    - ISO string ``YYYY-MM-DD``
    """
    val = ws.cell(row=row, column=col).value
    if val is None:
        return None

    # Already a date/datetime
    if isinstance(val, datetime):
        return val.date()
    if isinstance(val, date):
        return val

    # Excel serial number
    if isinstance(val, (int, float)):
        try:
            # openpyxl uses 1899-12-30 as epoch for date serial numbers
            from datetime import timedelta

            return (datetime(1899, 12, 30) + timedelta(days=int(val))).date()
        except Exception:
            logger.warning(
                "Cannot parse Excel serial date %r at %s%d",
                val,
                get_column_letter(col),
                row,
            )
            return None

    # String — try ISO format first, then common formats
    s = str(val).strip()
    formats = [
        "%Y-%m-%d",
        "%d/%m/%Y",
        "%m/%d/%Y",
        "%Y/%m/%d",
        "%d-%m-%Y",
        "%m-%d-%Y",
    ]
    for fmt in formats:
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue

    logger.warning("Cannot parse date %r at %s%d", s, get_column_letter(col), row)
    return None


def _cell_time_minutes(ws: Worksheet, row: int, col: int) -> int | None:
    """Read a cell value as minutes from midnight, or None.

    Handles:
    - Python ``time`` / ``datetime`` objects
    - Excel serial time number (float in [0,1))
    - String ``HH:MM`` or ``HH:MM:SS``
    """
    val = ws.cell(row=row, column=col).value
    if val is None:
        return None

    # Already a time/datetime
    if isinstance(val, time):
        return val.hour * 60 + val.minute
    if isinstance(val, datetime):
        return val.hour * 60 + val.minute

    # Excel serial time (fraction of day)
    if isinstance(val, float) and 0 <= val < 1:
        total_minutes = int(val * 24 * 60)
        return total_minutes

    # String
    s = str(val).strip()
    for fmt in ("%H:%M:%S", "%H:%M", "%I:%M %p", "%I:%M:%S %p"):
        try:
            t = datetime.strptime(s, fmt).time()
            return t.hour * 60 + t.minute
        except ValueError:
            continue

    logger.warning("Cannot parse time %r at %s%d", s, get_column_letter(col), row)
    return None


def _cell_bool(ws: Worksheet, row: int, col: int) -> bool:
    """Read a cell value as a boolean (truthy/falsy)."""
    val = ws.cell(row=row, column=col).value
    if val is None:
        return False
    if isinstance(val, bool):
        return val
    s = str(val).strip().lower()
    return s in ("yes", "true", "1", "y", "oui")
