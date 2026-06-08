"""Admin import routes for uploading Excel schedules.

Provides:
- ``POST /admin/import/preview`` — parse Excel and return preview JSON
- ``GET  /admin/import/teams/{team_id}/template`` — download blank template
- ``POST /admin/imports`` — create a persisted import record
- ``GET  /admin/imports`` — list imports for current user
- ``GET  /admin/imports/{import_id}`` — get full import record
- ``PUT  /admin/imports/{import_id}`` — update import record
- ``DELETE /admin/imports/{import_id}`` — delete import record
- ``GET  /admin/teams/{team_id}/merge-targets`` — get merge targets + auto-match
- ``POST /admin/imports/{import_id}/merge`` — execute merge into team
"""

from io import BytesIO
from typing import List

from fastapi import (
    APIRouter,
    Body,
    Depends,
    File,
    HTTPException,
    Query,
    UploadFile,
)
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from shared.database.database_collections import DatabaseCollections
from shared.logger import log_info
from shared.schemas.dto.import_merge import (
    MergeRequest,
    MergeResult,
    MergeTargetsResponse,
)
from shared.schemas.dto.import_preview import ImportPreviewDTO
from shared.schemas.dto.import_record import (
    CreateImportRequest,
    ImportRecordDTO,
    ImportRecordSummaryDTO,
    UpdateImportRequest,
)

from src.dependencies import get_db_collections, get_user_context
from src.dependencies.cerbos_authz_dependencies import get_cerbos_authz_service
from src.dependencies.import_merge_service import get_import_merge_service
from src.dependencies.import_persistence_service import (
    get_import_persistence_service,
)
from src.dependencies.import_service import get_import_service
from src.errors import NotAuthorizedError, handle_routes_errors
from src.integrations.authorization.cerbos_authz_service import (
    CerbosAuthzService,
)
from src.security.user_context import UserContext
from src.services.import_merge_service import ImportMergeService
from src.services.import_persistence_service import ImportPersistenceService
from src.services.import_service import ImportService

router = APIRouter()


# ── Preview endpoint ──────────────────────────────────────────────────────────


@router.post(
    "/admin/import/preview",
    status_code=201,
    response_model=ImportPreviewDTO,
)
async def preview_import(
    file: UploadFile = File(...),
    user_context: UserContext = Depends(get_user_context),
    import_service: ImportService = Depends(get_import_service),
    authz: CerbosAuthzService = Depends(get_cerbos_authz_service),
) -> ImportPreviewDTO:
    """Parse an Excel import file and return a preview of all extracted entities.

    No data is persisted — this is a read-only preview.  The caller must
    be a super admin (checked against the ``admin`` Cerbos resource).
    No team context is needed for preview. Team selection happens at
    the confirm/write phase.
    """
    response: ImportPreviewDTO
    try:
        # AuthZ: super-admin only
        if not await authz.check(
            user_context.user_id, "preview-import", "admin", "admin"
        ):
            raise NotAuthorizedError(
                "You do not have permission to import schedules"
            )

        # Validate file type
        if not file.filename or not (
            file.filename.endswith(".xlsx") or file.filename.endswith(".xls")
        ):
            raise ValueError("File must be an Excel workbook (.xlsx or .xls)")

        contents = await file.read()
        if not contents:
            raise ValueError("Uploaded file is empty")

        response = import_service.preview_import(contents)

    except Exception as e:
        log_info(f"Failed to preview import: {e}")
        handle_routes_errors(e)

    return response


# ── Template download endpoint ───────────────────────────────────────────────


@router.get("/admin/import/teams/{team_id}/template")
async def download_import_template(
    team_id: str,
    user_context: UserContext = Depends(get_user_context),
    authz: CerbosAuthzService = Depends(get_cerbos_authz_service),
    db_collections: DatabaseCollections = Depends(get_db_collections),
) -> StreamingResponse:
    """Download a blank Excel template for schedule import.

    The template contains three sheets with header rows:
    - ``members``
    - ``shifts``
    - ``schedule`` (header row only, no data)
    """
    try:
        if not await authz.check(
            user_context.user_id, "download-template", "admin", "admin"
        ):
            raise NotAuthorizedError(
                "You do not have permission to download import templates"
            )

        # Verify the target team exists
        team = db_collections.team_db.get_team_by_id(team_id)
        if team is None:
            raise ValueError(f"Team with id '{team_id}' does not exist")

        wb = _build_template_workbook()
        stream = BytesIO()
        wb.save(stream)
        stream.seek(0)

        response = StreamingResponse(
            stream,
            media_type=(
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            ),
        )
        response.headers["Content-Disposition"] = (
            f"attachment; filename=import_template_{team_id}.xlsx"
        )
        return response

    except Exception as e:
        log_info(f"Failed to generate import template: {e}")
        handle_routes_errors(e)


# ── Helpers ───────────────────────────────────────────────────────────────────


def _resolve_user_name(user_id: str, db: DatabaseCollections) -> str:
    """Resolve a user ID to a display name (first + last, or email fallback)."""
    user = db.user_db.get_user_by_id(user_id)
    if user is None:
        return user_id
    name = f"{user.first_name} {user.last_name}".strip()
    return name or user.email or user_id


# ── Import CRUD endpoints ────────────────────────────────────────────────────


@router.post(
    "/admin/imports",
    status_code=201,
    response_model=ImportRecordDTO,
)
async def create_import(
    req: CreateImportRequest = Body(...),
    user_context: UserContext = Depends(get_user_context),
    persistence_service: ImportPersistenceService = Depends(
        get_import_persistence_service
    ),
    authz: CerbosAuthzService = Depends(get_cerbos_authz_service),
    db_collections: DatabaseCollections = Depends(get_db_collections),
) -> ImportRecordDTO:
    """Create a new import record from preview data."""
    try:
        if not await authz.check(
            user_context.user_id, "manage-imports", "admin", "admin"
        ):
            raise NotAuthorizedError(
                "You do not have permission to manage imports"
            )

        record = persistence_service.create_import(
            preview_data=req.previewData,
            name=req.name,
            filename=req.filename,
            user_id=user_context.effective_user_id,
            team_id=req.teamId,
        )
        return record.to_dto(
            created_by_name=_resolve_user_name(
                record.created_by, db_collections
            )
        )

    except Exception as e:
        log_info(f"Failed to create import: {e}")
        handle_routes_errors(e)


@router.get(
    "/admin/imports",
    response_model=List[ImportRecordSummaryDTO],
)
async def list_imports(
    user_context: UserContext = Depends(get_user_context),
    persistence_service: ImportPersistenceService = Depends(
        get_import_persistence_service
    ),
    authz: CerbosAuthzService = Depends(get_cerbos_authz_service),
    db_collections: DatabaseCollections = Depends(get_db_collections),
) -> List[ImportRecordSummaryDTO]:
    """List all imports for the current user."""
    try:
        if not await authz.check(
            user_context.user_id, "manage-imports", "admin", "admin"
        ):
            raise NotAuthorizedError(
                "You do not have permission to view imports"
            )

        records = persistence_service.get_imports(
            user_id=user_context.effective_user_id
        )
        return [
            ImportRecordSummaryDTO(
                id=r.id,
                name=r.name,
                createdAt=r.created_at.timestamp(),
                updatedAt=r.updated_at.timestamp(),
                createdBy=r.created_by,
                createdByName=_resolve_user_name(r.created_by, db_collections),
                filename=r.filename,
                teamId=r.team_id,
                memberCount=len(r.members),
                shiftCount=len(r.shifts),
                requestCount=len(r.requests),
                assignmentCount=len(r.assignments),
            )
            for r in records
        ]

    except Exception as e:
        log_info(f"Failed to list imports: {e}")
        handle_routes_errors(e)


@router.get(
    "/admin/imports/{import_id}",
    response_model=ImportRecordDTO,
)
async def get_import(
    import_id: str,
    user_context: UserContext = Depends(get_user_context),
    persistence_service: ImportPersistenceService = Depends(
        get_import_persistence_service
    ),
    authz: CerbosAuthzService = Depends(get_cerbos_authz_service),
    db_collections: DatabaseCollections = Depends(get_db_collections),
) -> ImportRecordDTO:
    """Get a full import record by ID."""
    try:
        if not await authz.check(
            user_context.user_id, "manage-imports", "admin", "admin"
        ):
            raise NotAuthorizedError(
                "You do not have permission to view imports"
            )

        record = persistence_service.get_import(import_id)
        if record is None:
            raise HTTPException(status_code=404, detail="Import not found")
        return record.to_dto(
            created_by_name=_resolve_user_name(
                record.created_by, db_collections
            )
        )

    except HTTPException:
        raise
    except Exception as e:
        log_info(f"Failed to get import: {e}")
        handle_routes_errors(e)


@router.put(
    "/admin/imports/{import_id}",
    response_model=ImportRecordDTO,
)
async def update_import(
    import_id: str,
    req: UpdateImportRequest = Body(...),
    user_context: UserContext = Depends(get_user_context),
    persistence_service: ImportPersistenceService = Depends(
        get_import_persistence_service
    ),
    authz: CerbosAuthzService = Depends(get_cerbos_authz_service),
    db_collections: DatabaseCollections = Depends(get_db_collections),
) -> ImportRecordDTO:
    """Update an import record (partial update)."""
    try:
        if not await authz.check(
            user_context.user_id, "manage-imports", "admin", "admin"
        ):
            raise NotAuthorizedError(
                "You do not have permission to modify imports"
            )

        record = persistence_service.update_import(
            import_id=import_id,
            name=req.name,
            team_id=req.teamId,
            members=req.members,
            shifts=req.shifts,
            requests=req.requests,
            assignments=req.assignments,
        )
        return record.to_dto(
            created_by_name=_resolve_user_name(
                record.created_by, db_collections
            )
        )

    except ValueError as e:
        raise HTTPException(status_code=404, detail=str(e))
    except Exception as e:
        log_info(f"Failed to update import: {e}")
        handle_routes_errors(e)


# ── Merge endpoints ──────────────────────────────────────────────────────────


@router.get(
    "/admin/teams/{team_id}/merge-targets",
    response_model=MergeTargetsResponse,
)
async def get_merge_targets(
    team_id: str,
    import_id: str = Query(..., description="Import record ID"),
    user_context: UserContext = Depends(get_user_context),
    merge_service: ImportMergeService = Depends(get_import_merge_service),
    persistence_service: ImportPersistenceService = Depends(
        get_import_persistence_service
    ),
    authz: CerbosAuthzService = Depends(get_cerbos_authz_service),
) -> MergeTargetsResponse:
    """Get existing team workers/shifts plus auto-match suggestions.

    Returns the list of existing entities available as merge targets,
    along with pre-computed suggested mappings based on name/acronym
    matching."""
    try:
        if not await authz.check(
            user_context.user_id, "resolve-merge-targets", "admin", "admin"
        ):
            raise NotAuthorizedError(
                "You do not have permission to resolve merge targets"
            )

        record = persistence_service.get_import(import_id)
        if record is None:
            raise HTTPException(status_code=404, detail="Import not found")

        return merge_service.resolve_merge_targets(team_id, record)

    except HTTPException:
        raise
    except Exception as e:
        log_info(f"Failed to resolve merge targets: {e}")
        handle_routes_errors(e)


@router.post(
    "/admin/imports/{import_id}/merge",
    response_model=MergeResult,
)
async def execute_merge(
    import_id: str,
    req: MergeRequest = Body(...),
    user_context: UserContext = Depends(get_user_context),
    merge_service: ImportMergeService = Depends(get_import_merge_service),
    authz: CerbosAuthzService = Depends(get_cerbos_authz_service),
) -> MergeResult:
    """Execute the merge of imported data into the target team.

    Creates/updates workers, shifts, requests, and assignments
    according to the provided mappings.  Cascading exclusions are
    enforced: skipping a worker excludes its related requests and
    assignments."""
    try:
        if not await authz.check(
            user_context.user_id, "merge-imports", "admin", "admin"
        ):
            raise NotAuthorizedError(
                "You do not have permission to merge imports"
            )

        return merge_service.execute_merge(import_id, req)

    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
    except Exception as e:
        log_info(f"Failed to execute merge: {e}")
        handle_routes_errors(e)


@router.delete("/admin/imports/{import_id}", status_code=204)
async def delete_import(
    import_id: str,
    user_context: UserContext = Depends(get_user_context),
    persistence_service: ImportPersistenceService = Depends(
        get_import_persistence_service
    ),
    authz: CerbosAuthzService = Depends(get_cerbos_authz_service),
) -> None:
    """Delete an import record."""
    try:
        if not await authz.check(
            user_context.user_id, "manage-imports", "admin", "admin"
        ):
            raise NotAuthorizedError(
                "You do not have permission to delete imports"
            )

        persistence_service.delete_import(import_id)

    except ValueError as e:
        raise HTTPException(status_code=404, detail=str(e))
    except Exception as e:
        log_info(f"Failed to delete import: {e}")
        handle_routes_errors(e)


# ── Template builder ─────────────────────────────────────────────────────────


def _build_template_workbook() -> Workbook:
    """Build a blank template workbook with three sheets and header rows."""
    wb = Workbook()

    # ── Members sheet ─────────────────────────────────────────────────────
    ws_members = wb.active
    if ws_members is None:
        ws_members = wb.create_sheet("members")
    ws_members.title = "members"
    _write_headers(
        ws_members,
        [
            "name",
            "code",
            "start",
            "end",
            "skills",
            "contract",
            "desired",
            "duty/month",
            "leave",
        ],
    )
    # Add example row
    _write_row(
        ws_members,
        2,
        [
            "Dr. Example",
            "EX",
            "2026-01-01",
            "",
            "Cardiology;Surgery",
            "40",
            "40",
            "4",
            "20",
        ],
    )

    # ── Shifts sheet ──────────────────────────────────────────────────────
    ws_shifts = wb.create_sheet("shifts")
    _write_headers(
        ws_shifts,
        [
            "name",
            "code",
            "duty",
            "mandatory_rest",
            "start",
            "end",
            "staffing",
        ],
    )
    _write_row(
        ws_shifts,
        2,
        [
            "Morning",
            "AM",
            "",
            "",
            "08:00",
            "16:00",
            "",
        ],
    )
    _write_row(
        ws_shifts,
        3,
        [
            "Night Duty",
            "ND",
            "yes",
            "yes",
            "20:00",
            "08:00",
            "",
        ],
    )

    # ── Schedule sheet ────────────────────────────────────────────────────
    ws_schedule = wb.create_sheet("schedule")
    ws_schedule.cell(row=1, column=1, value="Worker").font = _bold()
    # A few example date columns
    from datetime import date, timedelta

    today = date.today()
    for i in range(14):
        cell = ws_schedule.cell(
            row=1, column=2 + i, value=(today + timedelta(days=i))
        )
        cell.font = _bold()
        cell.number_format = "YYYY-MM-DD"

    return wb


def _write_headers(ws, headers: list) -> None:
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = _bold()


def _write_row(ws, row_idx: int, values: list) -> None:
    for col_idx, value in enumerate(values, start=1):
        ws.cell(row=row_idx, column=col_idx, value=value)


def _bold():
    from openpyxl.styles import Font

    return Font(bold=True)
